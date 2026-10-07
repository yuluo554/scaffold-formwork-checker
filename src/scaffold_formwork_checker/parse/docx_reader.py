# -*- coding: utf-8 -*-
"""M3 docx 读取器：python-docx 抽取段落与表格（文本路线零识图）。

返回统一文档模型 {"paragraphs": [{"style","text"}], "tables": [[row][cell]]}，
供 extract.py 抽取管线消费。表格图片/扫描件不可读属文本路线固有局限
（README 已登记限制；品茗规范库表值转录走 engine/tables.py 人工双读通道）。
"""

from .card import ParseError


def read_docx(path):
    """读取 docx → 段落列表 + 表格三维列表（表→行→单元格文本）。"""
    try:
        from docx import Document
    except ImportError as exc:
        raise ParseError(
            "python-docx 未安装：请先 pip install \".[parse]\" 或 \".[dev]\"") from exc
    try:
        document = Document(path)
    except Exception as exc:
        raise ParseError("docx 文件无法读取：%s" % exc) from exc
    paragraphs = []
    for p in document.paragraphs:
        text = (p.text or "").strip()
        if not text:
            continue
        style = ""
        try:
            style = p.style.name if p.style is not None else ""
        except Exception:
            pass
        paragraphs.append({"style": style, "text": text})
    tables = []
    for table in document.tables:
        rows = []
        for row in table.rows:
            rows.append([(cell.text or "").strip() for cell in row.cells])
        tables.append(rows)
    return {"paragraphs": paragraphs, "tables": tables}
