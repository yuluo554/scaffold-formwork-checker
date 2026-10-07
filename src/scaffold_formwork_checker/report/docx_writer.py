# -*- coding: utf-8 -*-
"""M4 报告导出公共件：docx 无时间戳落盘 + 共享常量（plan/04 §7）。

确定性纪律：
- 内容字段无时间戳（封面/签署栏日期一律手填，不写 datetime.now()）；
- 落盘前重写 zip：全条目 date_time=(1980,1,1,0,0,0)，core.xml/app.xml 覆写
  固定内容（做法同 synth/generator.py `_normalize_zip`）→ 同输入位级一致；
- 零外链：不添加任何超链接（DoD 断言测试扫描 TargetMode="External" 锁死）。
"""

import io
import os
import zipfile

from .. import __version__

_ZIP_DATE = (1980, 1, 1, 0, 0, 0)

_FIXED_CORE_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
    '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
    'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" '
    'xmlns:dcmitype="http://purl.org/dc/dcmitype/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
    '<dc:title>sfc-report</dc:title><dc:creator>sfc-report</dc:creator>'
    '<cp:lastModifiedBy>sfc-report</cp:lastModifiedBy>'
    '<dcterms:created xsi:type="dcterms:W3CDTF">2026-01-01T00:00:00Z</dcterms:created>'
    '<dcterms:modified xsi:type="dcterms:W3CDTF">2026-01-01T00:00:00Z</dcterms:modified>'
    '</cp:coreProperties>'
).encode("utf-8")
_FIXED_APP_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
    '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
    'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVT">'
    '<Application>Microsoft Office Word</Application><AppVersion>16.0000</AppVersion>'
    '</Properties>'
).encode("utf-8")

# 免责声明（强制节，验算书/核查报告共用；口径与 README 免责声明一致）
DISCLAIMER = (
    "免责声明：本文件由 scaffold-formwork-checker（sfc）v%s 自动生成，定位为辅助"
    "验算与分级工具，不替代专项施工方案编制、论证与审批。验算与判定结果必须由"
    "具备相应资格的专业人员复核后方可使用；本文件载明的一切结论均不构成工程决策"
    "依据。依据规范：JGJ 130-2011、JGJ 162-2008、住建部令第37号（2019修正）、"
    "建办质〔2018〕31号。" % __version__
)

TOOL_LINE = "生成工具：scaffold-formwork-checker（sfc）v%s（自动生成，零外链、无时间戳）" % __version__


def normalize_docx_bytes(raw):
    """重写 docx zip：固定条目时间戳、覆写 core/app 元数据（无时间戳纪律）。"""
    src = zipfile.ZipFile(io.BytesIO(raw))
    try:
        infos = src.infolist()
        payloads = [(i, src.read(i.filename)) for i in infos]
    finally:
        src.close()
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in payloads:
            zi = zipfile.ZipInfo(info.filename, date_time=_ZIP_DATE)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = info.external_attr
            if info.filename == "docProps/core.xml":
                data = _FIXED_CORE_XML
            elif info.filename == "docProps/app.xml":
                data = _FIXED_APP_XML
            z.writestr(zi, data)
    return out.getvalue()


def save_normalized(doc, out_path):
    """python-docx 文档 → 规范化 docx 字节落盘（父目录不存在则创建）。"""
    buf = io.BytesIO()
    doc.save(buf)
    data = normalize_docx_bytes(buf.getvalue())
    parent = os.path.dirname(os.path.abspath(out_path))
    if parent and not os.path.isdir(parent):
        os.makedirs(parent)
    with open(out_path, "wb") as f:
        f.write(data)
    return out_path


def new_document(base_font="SimSun", east_asia="宋体", size_pt=12):
    """统一版式的新文档：正文宋体小四（做法同 synth 生成器，设置失败不致命）。"""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = base_font
    style.font.size = Pt(size_pt)
    try:
        from docx.oxml.ns import qn
        style.element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    except Exception:
        pass
    return doc


def fmt_num(x):
    """数值定形（两位小数去尾零），与 synth/consistency 的展示口径一致。"""
    s = "%.2f" % float(x)
    s = s.rstrip("0").rstrip(".")
    return s if s else "0"


def fmt_num4(x):
    """数值定形（四位小数去尾零）：比值/中间量展示用（0.2814 不丢精度）。"""
    s = "%.4f" % float(x)
    s = s.rstrip("0").rstrip(".")
    return s if s else "0"


def fmt_value(value):
    """任意参数卡取值 → 展示文本（list 用分号连接，bool 用是/否）。"""
    if isinstance(value, list):
        return "；".join(fmt_value(v) for v in value)
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, float):
        s = "%.4f" % value
        s = s.rstrip("0").rstrip(".")
        return s if s else "0"
    if value is None:
        return "—"
    return str(value)


def add_table(doc, headers, rows):
    """加带表头的网格表；rows 为字符串两维列表（定列数=表头数，缺位补空）。"""
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i in range(len(headers)):
            cells[i].text = row[i] if i < len(row) else ""
    return table


__all__ = [
    "DISCLAIMER", "TOOL_LINE", "add_table", "fmt_num", "fmt_value",
    "new_document", "normalize_docx_bytes", "save_normalized",
]
