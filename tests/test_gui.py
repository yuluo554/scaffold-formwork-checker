# -*- coding: utf-8 -*-
"""M5 GUI offscreen 测试（plan/05 M5 DoD ①）。

纪律（skill 实录三条坑）：
- QT_QPA_PLATFORM=offscreen 必须在首次导入 PySide6 之前设置（本文件在
  import PySide6 前设置）；
- 模态 QMessageBox 会挂死 offscreen 批——页签代码零模态框，通知走
  可注入 notify 信号；文件对话框只出现在按钮槽内，测试 monkeypatch 屏蔽；
- 测试不 show() 上屏（offscreen grab 无系统字体，截图另走 native QPA）。

PySide6 未安装时整模块声明式跳过（CI 不装 gui extras——skip 数在发布期
收集数对账中逐项归因）。
"""
import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")  # 必须先于 PySide6 导入

from pathlib import Path

import pytest

pytest.importorskip("PySide6")
pytest.importorskip("docx")

from PySide6.QtWidgets import QApplication, QFileDialog

from scaffold_formwork_checker.engine import CardError
from scaffold_formwork_checker.gui.app import TAB_ORDER, MainWindow
from scaffold_formwork_checker.gui.bench_tab import BenchTab
from scaffold_formwork_checker.gui.calc_tab import CalcTab
from scaffold_formwork_checker.gui.check_tab import CheckTab
from scaffold_formwork_checker.gui.consistency_tab import ConsistencyTab
from scaffold_formwork_checker.gui.grading_tab import GradingTab
from scaffold_formwork_checker.gui.pipelines import (
    run_calc_pipeline, run_check_pipeline, run_consistency_pipeline,
)

_REPO = Path(__file__).resolve().parent.parent
_SYNTH = _REPO / "data" / "synthetic"


@pytest.fixture(scope="session")
def qapp():
    app = QApplication.instance() or QApplication([])
    yield app


@pytest.fixture()
def win(qapp):
    window = MainWindow()
    yield window
    window.deleteLater()


def _notifies(tab):
    """接住 notify 信号的通道（测试内可注入断言）。"""
    box = []
    tab.notify.connect(box.append)
    return box


# ---- 主窗体 ----

def test_main_window_five_tabs(qapp):
    window = MainWindow()
    titles = [window.tabs.tabText(i) for i in range(window.tabs.count())]
    assert titles == list(TAB_ORDER)
    window.deleteLater()


def test_main_window_tabs_share_data_dir_getter(win):
    data_dir = win.calc_tab._get_data_dir()
    assert data_dir and (Path(data_dir) / "knowledge" / "clauses").is_dir()


# ---- 验算页 ----

def test_calc_default_cards_run_all_modules(qapp, win):
    """默认卡（转录自算例）逐模块直接跑通引擎——字段规格与引擎校验无漂移。"""
    tab = win.calc_tab
    for mid in ("M-1", "M-2", "M-3", "M-4", "M-5", "M-6"):
        idx = tab.module_combo.findData(mid)
        tab.module_combo.setCurrentIndex(idx)
        result = tab.run_calc()
        assert result["status"] == "ok", (mid, result.get("status_note"))
        assert result["checks"]


def test_calc_bad_number_notifies_not_modal(qapp):
    tab = CalcTab(lambda: None)
    notes = _notifies(tab)
    tab._fields["step_m"].setText("abc")
    tab._on_run()
    assert notes and "须为数值" in notes[0]


def test_calc_missing_data_dir_notifies(qapp, tmp_path):
    """显式指定无效数据目录 → KnowledgeError（不静默回退口径的 GUI 呈现）。"""
    tab = CalcTab(lambda: str(tmp_path))
    notes = _notifies(tab)
    tab._on_run()
    assert notes and "未找到数据目录" in notes[0]


def test_calc_show_result_renders_table(qapp, win):
    tab = win.calc_tab
    tab._on_run()
    assert tab.table.rowCount() > 0
    assert tab.export_button.isEnabled()


def test_calc_pipeline_detect_module_matches_engine(qapp):
    card = {"category": "formwork_support", "slab_thickness_m": 0.12,
            "step_m": 1.5, "pole_spacing_m": [1.0, 1.0],
            "loads": {"g1k_kN_per_m2": 0.75, "concrete_kN_per_m3": 24.0,
                      "q1k_kN_per_m2": 1.0}}
    result = run_calc_pipeline(card)
    assert result["module_id"] == "M-6"
    assert result["status"] == "ok"


# ---- 方案核查页 ----

def test_check_tab_two_sheets_finds_violation(qapp, win):
    tab = win.check_tab
    tab.set_scheme_path(str(_SYNTH / "SY-two-sheets-01.docx"))
    notes = _notifies(tab)
    tab._on_run()
    rule_ids = [tab.table.item(i, 0).text() for i in range(tab.table.rowCount())]
    assert "R-consistency-step" in rule_ids
    assert any("违规" for i in range(tab.table.rowCount())
               if tab.table.item(i, 2).text() == "违规")
    assert notes and "核查完成" in notes[-1]


def test_check_tab_unknown_slots_displayed(qapp, win):
    """低置信度槽位显式列出（宁缺勿错口径的 GUI 呈现）。"""
    tab = win.check_tab
    tab.set_scheme_path(str(_SYNTH / "SY-grading-missing-01.docx"))
    tab._on_run()
    report = tab.report
    if report["scheme_card"]["unknown_slots"]:
        assert tab.slots_label.isVisibleTo(tab)
        assert "待人工确认槽位" in tab.slots_label.text()


def test_check_tab_missing_file_notifies(qapp, win):
    tab = win.check_tab
    tab.set_scheme_path(str(_REPO / "output" / "no-such-file.docx"))
    notes = _notifies(tab)
    tab._on_run()
    assert notes and "核查未执行" in notes[0]


# ---- 一致性页 ----

def test_consistency_tab_same_document_diff(qapp, win):
    tab = win.consistency_tab
    tab.set_paths(str(_SYNTH / "SY-two-sheets-01.docx"))
    tab._on_run()
    params = [tab.table.item(i, 1).text() for i in range(tab.table.rowCount())]
    assert "step" in params


def test_consistency_pipeline_two_files(qapp):
    """方案文件 + 验算书文件分立输入：取 A 卡方案、B 卡验算书。"""
    outcome = run_consistency_pipeline(
        str(_SYNTH / "SY-two-sheets-01.docx"), str(_SYNTH / "SY-clean-01.docx"))
    assert outcome["scheme_doc"] and outcome["calcbook_doc"]
    assert outcome["findings"] or outcome["confirmations"]


# ---- 危大分级页 ----

def test_grading_tab_definite(qapp, win):
    tab = win.grading_tab
    tab.set_scheme_path(str(_SYNTH / "SY-clean-01.docx"))
    tab._on_run()
    assert tab.result["status"] == "definite"
    assert tab.result["level"] in ("none", "weida", "chaoguimo")
    assert "结论：" in tab.summary_label.text()


def test_grading_tab_pending_shows_unknown(qapp, win, tmp_path):
    """缺判定参数 → pending（待人工确认，不硬判）的 GUI 呈现。

    冻结 fixtures 全部判 definite（见 test_grading.py），pending 语义用
    "有类型关键词但无搭设高度"的最小 docx 现做（dev 依赖 python-docx）。
    """
    from docx import Document
    doc = Document()
    doc.add_heading("测试工程落地式扣件钢管脚手架专项施工方案", level=0)
    doc.add_paragraph("本工程采用落地式扣件钢管脚手架（双排），逐层搭设。")
    path = tmp_path / "SY-gui-pending.docx"
    doc.save(str(path))
    tab = win.grading_tab
    tab.set_scheme_path(str(path))
    tab._on_run()
    assert tab.result["status"] == "pending"
    assert "build_height" in tab.result.get("unknown_params", [])
    assert "待人工确认" in tab.summary_label.text()


def test_grading_tab_param_overrides(qapp, win):
    tab = win.grading_tab
    tab.set_scheme_path(str(_SYNTH / "SY-clean-01.docx"))
    tab.param_edit.setText("build_height=56")
    tab._on_run()
    assert tab.result["status"] == "definite"
    assert tab.result["level"] in ("weida", "chaoguimo")


def test_grading_bad_override_notifies(qapp, win):
    tab = win.grading_tab
    tab.set_scheme_path(str(_SYNTH / "SY-clean-01.docx"))
    tab.param_edit.setText("bad-form")
    notes = _notifies(tab)
    tab._on_run()
    assert notes and "name=value" in notes[0]


# ---- 基准页 ----

def test_bench_tab_all_pass(qapp, win):
    tab = win.bench_tab
    notes = _notifies(tab)
    tab._on_run()
    assert tab.outcome["all_pass"] is True
    assert tab.table.rowCount() == 4
    verdicts = [tab.table.item(i, 4).text() for i in range(tab.table.rowCount())]
    assert verdicts == ["达标"] * 4
    assert notes and "全部达标" in notes[-1]


# ---- 导出（文件对话框 monkeypatch 屏蔽，不触模态）----

def test_export_calc_report(qapp, win, tmp_path, monkeypatch):
    tab = win.calc_tab
    tab._on_run()
    out = tmp_path / "验算书.docx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName",
                        staticmethod(lambda *a, **k: (str(out), "")))
    tab._on_export()
    assert out.is_file() and out.stat().st_size > 0


def test_export_check_report_no_external_links(qapp, win, tmp_path, monkeypatch):
    from tests.report_helpers import assert_no_external_links
    tab = win.check_tab
    tab.set_scheme_path(str(_SYNTH / "SY-two-sheets-01.docx"))
    tab._on_run()
    out = tmp_path / "核查报告.docx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName",
                        staticmethod(lambda *a, **k: (str(out), "")))
    tab._on_export()
    assert out.is_file()
    assert_no_external_links(str(out))


def test_export_cancelled_is_noop(qapp, win, tmp_path, monkeypatch):
    tab = win.calc_tab
    tab._on_run()
    monkeypatch.setattr(QFileDialog, "getSaveFileName",
                        staticmethod(lambda *a, **k: ("", "")))
    notes = _notifies(tab)
    tab._on_export()
    assert not list(tmp_path.iterdir())
    assert not notes  # 取消导出不打扰


# ---- 探针 ----

def test_probe_returns_zero(qapp):
    from scaffold_formwork_checker.gui import app as gui_app
    assert gui_app.main(["--probe"]) == 0
