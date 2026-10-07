# -*- coding: utf-8 -*-
"""M4 报告守门：docx 核查报告（M4 DoD ②③ + 报告内容纪律）。

- 位级一致：同输入两次生成字节完全相同；
- 0 外链：全部 .rels 条目无 TargetMode="External" 超链关系；
- 全节结构：概要/参数卡/逐项核查/违规清单/分级结论/待人工确认/免责声明；
- 违规项载级别+建议+条款+原文摘录（knowledge 回填）。
"""

import os
import zipfile

import pytest

from scaffold_formwork_checker.parse import parse_scheme
from scaffold_formwork_checker.report import DISCLAIMER, generate_check_report
from scaffold_formwork_checker.rules import load_checks_table, run_checks
from report_helpers import assert_no_external_links, assert_no_timestamps, read_plain_text

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(REPO_ROOT, "data")
SYNTH_DIR = os.path.join(DATA_DIR, "synthetic")


@pytest.fixture(scope="module")
def knowledge():
    from scaffold_formwork_checker.engine import load_knowledge
    return load_knowledge(DATA_DIR)


def _build_report(fixture, knowledge):
    path = os.path.join(SYNTH_DIR, fixture + ".docx")
    parsed = parse_scheme(path)
    report = run_checks(parsed.scheme_card, parsed.calcbook_card, knowledge,
                        load_checks_table())
    report = dict(report)
    report["file"] = path
    report["warnings"] = list(parsed.warnings)
    return report


def test_check_report_byte_identical_on_regeneration(tmp_path, knowledge):
    report = _build_report("SY-two-sheets-01", knowledge)
    out1 = str(tmp_path / "a.docx")
    out2 = str(tmp_path / "b.docx")
    generate_check_report(report, out1, knowledge=knowledge)
    generate_check_report(report, out2, knowledge=knowledge)
    with open(out1, "rb") as f:
        b1 = f.read()
    with open(out2, "rb") as f:
        b2 = f.read()
    assert b1 == b2 and len(b1) > 0


def test_check_report_zero_external_links_and_no_timestamps(tmp_path, knowledge):
    report = _build_report("SY-two-sheets-01", knowledge)
    out = str(tmp_path / "check.docx")
    generate_check_report(report, out, knowledge=knowledge)
    assert_no_external_links(out)
    assert_no_timestamps(out)


def test_check_report_clean_fixture_all_sections(tmp_path, knowledge):
    report = _build_report("SY-clean-01", knowledge)
    out = str(tmp_path / "check-clean.docx")
    generate_check_report(report, out, knowledge=knowledge)
    text = read_plain_text(out)
    for kw in ("专项施工方案核查报告", "一、核查概要", "二、方案概况与参数卡",
               "三、逐项核查结果", "四、违规清单", "五、危大分级结论",
               "六、待人工确认项", "免责声明"):
        assert kw in text, kw
    assert "本次核查无违规项。" in text
    assert "非危大" in text
    assert "（无待人工确认项）" in text
    assert DISCLAIMER in text


def test_check_report_violation_carries_level_advice_excerpt(tmp_path, knowledge):
    """两张皮违规项：级别/建议/条款/原文摘录逐项在报告内。"""
    report = _build_report("SY-two-sheets-01", knowledge)
    out = str(tmp_path / "check-ts.docx")
    generate_check_report(report, out, knowledge=knowledge)
    text = read_plain_text(out)
    assert "R-consistency-step" in text
    assert "不合格" in text
    assert "整改建议：统一方案正文与验算书取值" in text
    assert "37号令-第十六条" in text
    assert "依据摘录：37号令-第十六条《住建部令第37号》" in text
    assert "不得擅自修改专项施工方案" in text
    assert DISCLAIMER in text


def test_check_report_grading_chaoguimo_duty_and_basis(tmp_path, knowledge):
    report = _build_report("SY-grading-missing-01", knowledge)
    out = str(tmp_path / "check-grading.docx")
    generate_check_report(report, out, knowledge=knowledge)
    text = read_plain_text(out)
    assert "超过一定规模的危大工程" in text
    assert "义务提示：专项施工方案+专家论证（专家≥5名）" in text
    assert "建办质〔2018〕31号" in text
    assert "R-grading-expert-review" in text
    assert "严重" in text


def test_check_report_unknown_slots_listed(tmp_path, knowledge):
    """低置信度/未解析槽位须在报告显式列出（宁缺勿错可回溯）。

    冻结 docx fixtures 参数表齐全，此处用最小方案（无参数表）触发 unknown_slots。
    """
    from docx import Document

    doc = Document()
    doc.add_heading("某工程落地式扣件钢管脚手架专项施工方案", level=0)
    doc.add_paragraph("本工程采用落地式扣件钢管脚手架（双排）。")
    path = str(tmp_path / "minimal-scheme.docx")
    doc.save(path)
    parsed = parse_scheme(path)
    assert parsed.scheme_card.unknown_slots, "前置：最小方案应产生 unknown_slots"
    report = run_checks(parsed.scheme_card, parsed.calcbook_card, knowledge,
                        load_checks_table())
    out = str(tmp_path / "check-slots.docx")
    generate_check_report(report, out, knowledge=knowledge)
    text = read_plain_text(out)
    assert "未解析/低置信度槽位" in text
    assert "step" in text
