# -*- coding: utf-8 -*-
"""M4 CLI 守门：sfc report calc/check 退出码与输出结构（0/1/2 口径）。

报告已生成但内容含违规/待确认 → exit 1（降级完成）；输入不可用 → 2。
"""

import json
import os

import pytest

from scaffold_formwork_checker.cli import EXIT_DEGRADED, EXIT_INPUT_ERROR, EXIT_OK, main

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(REPO_ROOT, "data")
SYNTH_DIR = os.path.join(DATA_DIR, "synthetic")


@pytest.fixture()
def card_json(tmp_path):
    with open(os.path.join(DATA_DIR, "examples", "EX-lizhigan-nw-001.json"),
              encoding="utf-8") as f:
        card = json.load(f)["input_card"]
    path = tmp_path / "card.json"
    path.write_text(json.dumps(card, ensure_ascii=False), encoding="utf-8")
    return str(path)


def test_report_calc_ok_exit_zero(card_json, tmp_path, capsys):
    out = str(tmp_path / "calc.docx")
    code = main(["report", "calc", card_json, "-o", out])
    assert code == EXIT_OK
    summary = json.loads(capsys.readouterr().out)
    assert summary["output"] == out
    assert summary["status"] == "ok"
    assert os.path.isfile(out)


def test_report_calc_missing_card_exits_input_error(tmp_path):
    assert main(["report", "calc", "no-such.json",
                 "-o", str(tmp_path / "x.docx")]) == EXIT_INPUT_ERROR


def test_report_calc_bad_json_exits_input_error(tmp_path, capsys):
    bad = tmp_path / "bad.json"
    bad.write_text("{not json", encoding="utf-8")
    assert main(["report", "calc", str(bad),
                 "-o", str(tmp_path / "x.docx")]) == EXIT_INPUT_ERROR
    assert "JSON 解析失败" in capsys.readouterr().err


def test_report_calc_bad_data_dir_exits_input_error(card_json, tmp_path):
    assert main(["report", "calc", card_json, "-o", str(tmp_path / "x.docx"),
                 "--data-dir", "no-such-dir"]) == EXIT_INPUT_ERROR


def test_report_check_clean_exit_zero(tmp_path, capsys):
    out = str(tmp_path / "clean.docx")
    code = main(["report", "check", os.path.join(SYNTH_DIR, "SY-clean-01.docx"),
                 "-o", out])
    assert code == EXIT_OK
    summary = json.loads(capsys.readouterr().out)
    assert summary["violation"] == 0
    assert summary["confirmations"] == 0
    assert os.path.isfile(out)


def test_report_check_violation_exits_degraded(tmp_path, capsys):
    """报告已生成（产物在盘上），内容含违规 → exit 1。"""
    out = tmp_path / "ts.docx"
    code = main(["report", "check", os.path.join(SYNTH_DIR, "SY-two-sheets-01.docx"),
                 "-o", str(out)])
    assert code == EXIT_DEGRADED
    summary = json.loads(capsys.readouterr().out)
    assert summary["violation"] == 1
    assert out.is_file(), "降级完成时报告仍应已落盘"


def test_report_check_confirmation_exits_degraded(tmp_path):
    """参数缺失（无参数表的最小方案）→ 待人工确认 → exit 1，报告仍落盘。"""
    from docx import Document

    doc = Document()
    doc.add_heading("某工程落地式扣件钢管脚手架专项施工方案", level=0)
    doc.add_paragraph("本工程采用落地式扣件钢管脚手架（双排）。")
    minimal = tmp_path / "minimal-scheme.docx"
    doc.save(str(minimal))
    out = tmp_path / "p.docx"
    code = main(["report", "check", str(minimal), "-o", str(out)])
    assert code == EXIT_DEGRADED
    assert out.is_file(), "降级完成时报告仍应已落盘"


def test_report_check_missing_file_exits_input_error(tmp_path):
    assert main(["report", "check", "no-such.docx",
                 "-o", str(tmp_path / "x.docx")]) == EXIT_INPUT_ERROR


def test_report_check_garbage_file_exits_input_error(tmp_path, capsys):
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"not a docx")
    assert main(["report", "check", str(bad),
                 "-o", str(tmp_path / "x.docx")]) == EXIT_INPUT_ERROR
    assert "解析失败" in capsys.readouterr().err


def test_report_no_kind_exits_input_error(capsys):
    assert main(["report"]) == EXIT_INPUT_ERROR
    assert "calc|check" in capsys.readouterr().err
